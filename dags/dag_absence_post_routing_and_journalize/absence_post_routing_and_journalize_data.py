import logging
import re

from io import BytesIO
from openpyxl import load_workbook


logger = logging.getLogger(__name__)

DEPARTMENT_COLUMN = "NUV."
EMAIL_COLUMNS = ("Email 1", "Email 2", "Email 3", "Email 4")
CPR_REGEX = re.compile(
    r"(?<!\d)(?:"
    r"(?:0[1-9]|[12][0-9]|3[01])(?:0[13578]|10|12)\d{2}|"
    r"(?:0[1-9]|[12][0-9]|30)(?:0[469]|11)\d{2}|"
    r"(?:0[1-9]|1[0-9]|2[0-8])02\d{2}|"
    r"2902(?:00|[2468][048]|[13579][26])"
    r")(?:-?\d{4})(?!\d)"
)


def build_department_email_map(excel_bytes: bytes) -> dict[str, list[str]]:
    """
    Build a department-to-recipient mapping from an SD-ORG Excel file.

    :param excel_bytes: Excel file bytes containing department and email data.
    :return: Dictionary mapping department codes to recipient email addresses.
    """
    if not excel_bytes:
        raise ValueError("Excel attachment is empty")

    workbook = load_workbook(
        filename=BytesIO(excel_bytes),
        read_only=True,
        data_only=True,
    )
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        department_index = None
        email_indexes: list[int] = []
        header_row_number = 0
        for header_row_number, headers in enumerate(rows, start=1):
            normalized_headers = {
                str(value).strip().casefold(): index
                for index, value in enumerate(headers)
                if value is not None and str(value).strip()
            }
            department_index = normalized_headers.get(DEPARTMENT_COLUMN.casefold())
            email_indexes = [
                normalized_headers[column.casefold()]
                for column in EMAIL_COLUMNS
                if column.casefold() in normalized_headers
            ]
            if department_index is not None and email_indexes:
                break
        else:
            raise ValueError(f"Excel file must contain columns '{DEPARTMENT_COLUMN}' and at least one email column")

        department_email_map: dict[str, list[str]] = {}
        for row_number, row in enumerate(rows, start=header_row_number + 1):
            department = row[department_index] if department_index < len(row) else None
            emails = [
                str(row[index]).strip()
                for index in email_indexes
                if index < len(row) and row[index] is not None and str(row[index]).strip()
            ]

            if department is None and not emails:
                continue
            if not emails:
                continue
            if department is None or not str(department).strip():
                raise ValueError(f"Row {row_number} has no '{DEPARTMENT_COLUMN}' value")

            department_key = str(department).strip()
            recipients = department_email_map.setdefault(department_key, [])
            for recipient in emails:
                if recipient not in recipients:
                    recipients.append(recipient)

        if not department_email_map:
            raise ValueError("Excel file does not contain any department email mappings")

        return department_email_map
    finally:
        workbook.close()


def extract_cpr_from_pdf(pdf_bytes: bytes) -> str:
    """
    Extract one date-valid CPR number from a PDF as ten digits.

    :param pdf_bytes: PDF file bytes containing a CPR number.
    :return: The first valid CPR number found, without a separator.
    """
    if not pdf_bytes:
        raise ValueError("PDF attachment is empty")

    import fitz

    try:
        with fitz.open(stream=pdf_bytes, filetype="pdf") as document:
            pdf_text = "\n".join(page.get_text() for page in document)
    except (fitz.FileDataError, RuntimeError) as exc:
        raise ValueError("PDF attachment could not be parsed") from exc

    matches = CPR_REGEX.findall(pdf_text)
    if not matches:
        raise ValueError("No valid CPR number found in PDF")

    return matches[0].replace("-", "")
