import io
import re
import fitz
import pandas as pd
import logging
from openpyxl.utils import get_column_letter
from typing import Any, Iterable, Sequence
from airflow.providers.http.hooks.http import HttpHook
from rkdigi.email_handling import EmailReader
from datetime import datetime, timezone

logger = logging.getLogger(__name__)

INSUBIZ_EXCEL_FIELDS = [
    "registration",
    "registrationDate",
    "make",
    "model",
    "modelYear",
    "id",
    "chassisNumber",
    "cart",
    "endDate",
    "customer_name",
    "usageCategory_text",
    "usage_text",
    "Level1",
    "Level2",
    "Level3",
    "Level4",
    "Level5",
    "Level6",
]

INSUBIZ_COLUMN_RENAME_MAP = {
    "registration": "Reg.nr.",
    "registrationDate": "Reg.dato",
    "make": "Mærke",
    "model": "Model",
    "modelYear": "Årgang",
    "chassisNumber": "Stelnr.",
    "cart": "Træk",
    "endDate": "Afg.dato",
    "customer_name": "Kundenavn",
    "usageCategory_text": "Anvendelse",
    "usage_text": "Art",
}

MOTORSTYRELSEN_PDF_COLUMNS = [
    "Nummer (DPRN)",
    "Stelnummer (DKKØ)",
    "Køretøj art beskrivelse (DKRG)",
    "Mærke beskrivelse (DKKØ)",
    "Variant beskrivelse (DKKØ)",
    "Dato - første registrering (DKKØ)",
    "Model år (DKKØ)",
    "Bruger p-nr navn",
    "Model beskrivelse (DKKØ)"
]

INSUBIZ_ROOT_CUSTOMER_ID = 216146  # RK
INSUBIZ_MAX_LEVELS = 6

USAGE_TEXT_TO_CATEGORY_MAP = {
    "22": "Personbil",
    "34": "Varebil",
    "25": "Påhængsvogn/trailer",
    "0": "Ukendt",
    "30": "Traktor",
    "11": "Knallert",
    "1": "Arbejdsmaskine",
    "24": "Påhængsredskab",
    "21": "Motorredskab",
    "4": "Bus",
    "36": "Campingvogn",
    "7": "Græsklipper",
    "31": "Traktorpåhængsvogn",
    "15": "Lastbil",
    "41": "Stor personbil",
    "37": "Skurvogn/mandskabsvogn",
    "3": "Brandbil med pakninger",
}


def _format_datetime_column(df: pd.DataFrame, col: str) -> None:
    """
    If the column does not exist, the function returns without changes.

    :param df: DataFrame containing the column to format.
    :param col: Name of the column to parse and format as datetime.
    :return: None. The DataFrame is modified in place.
    """
    if col not in df.columns:
        return
    dt = pd.to_datetime(df[col], errors="coerce")
    df[col] = dt.dt.strftime("%Y-%m-%d %H:%M:%S")


def normalize_insubiz_df(df: pd.DataFrame) -> pd.DataFrame:
    """
    Renames source columns to the internal export naming and formats known
    datetime columns to a consistent timestamp string format.

    :param df: Raw Insubiz DataFrame.
    :return: Normalized DataFrame with renamed and formatted columns.
    """
    out = df.rename(columns=INSUBIZ_COLUMN_RENAME_MAP)

    for col in ("Reg.dato", "Afg.dato", "start_date"):
        _format_datetime_column(out, col)

    return out


def _get_insubiz_credentials(http_hook: HttpHook) -> tuple[str, str]:
    """
    Read Insubiz API credentials from the Airflow HTTP connection extras.

    Expects the Airflow connection (http_conn_id = http_hook.http_conn_id) to contain
    extras keys: ``api_key`` and ``secret_key``.

    :param http_hook: Airflow HttpHook configured for the Insubiz Cloud API connection.
    :return: (api_key, secret_key)
    """
    conn = http_hook.get_connection(http_hook.http_conn_id)

    extras = conn.extra_dejson

    api_key = extras.get("api_key")
    secret_key = extras.get("secret_key")

    if not api_key or not secret_key:
        raise ValueError(
            "Missing api_key or secret_key in Airflow connection extras."
        )

    return api_key, secret_key


def _insubiz_sign_in(http_hook: HttpHook) -> str:
    """
    Authenticate against the Insubiz Cloud API and return a bearer token (JWT).

    :param http_hook: Airflow HttpHook configured for the Insubiz Cloud API connection.
    :return: JWT token string to be used as ``Authorization: Bearer <token>``.
    """
    api_key, secret_key = _get_insubiz_credentials(http_hook=http_hook)

    session = http_hook.get_conn()

    session.auth = None

    payload = {
        "apiKey": api_key,
        "secretKey": secret_key,
    }

    response = session.post(
        url=f"{http_hook.base_url}/api/v1.3/Authentication/SignInAsync",
        json=payload,
        headers={
            "Content-Type": "application/json",
            "Accept": "application/json",
        },
    )

    response.raise_for_status()

    body: Any = response.json()

    if not isinstance(body, dict):
        raise ValueError(f"Unexpected SignIn response type: {type(body)}")

    if not body.get("isAuthenticated"):
        msg = body.get("message") or "Authentication failed"
        raise ValueError(f"Insubiz SignIn failed: {msg}")

    token = body.get("token")

    if not token or not isinstance(token, str):
        raise ValueError("Insubiz SignIn response missing 'token'")

    return token


def fetch_insubiz_vehicles(
    http_hook: HttpHook,
    page_size: int = 500,
) -> list[dict[str, Any]]:
    """
    Fetch all vehicles from the Insubiz Cloud API (paged)

    :param http_hook: Airflow HttpHook configured for the Insubiz Cloud API connection.
    :param page_size: Page size used for pagination requests.
    :return: List of vehicle data (concatenated across pages).
    """
    logger.info("Fetching vehicles from Insubiz Cloud API ...")

    jwt = _insubiz_sign_in(http_hook=http_hook)

    page_no = 1
    all_rows: list[dict] = []

    headers = {
        "Authorization": f"Bearer {jwt}",
        "Content-Type": "application/json",
        "Accept": "application/json",
    }

    while True:
        payload = {
            "pageNo": page_no,
            "pageSize": page_size,
        }

        response = http_hook.run(
            endpoint="/api/v1.3/Vehicle/GetVehiclesPagedAsync",
            json=payload,
            headers=headers,
        )

        response.raise_for_status()

        body: Any = response.json()

        if isinstance(body, dict):
            batch = body.get("data", []) or []
        elif isinstance(body, list):
            batch = body
        else:
            raise ValueError(f"Unexpected response type from vehicles endpoint: {type(body)}")

        if not isinstance(batch, list):
            raise ValueError("Expected a list of vehicles in response under key 'data'.")

        all_rows.extend(batch)

        logger.info(f"Fetched pageNo={page_no} rowws={len(batch)} total={len(all_rows)}")

        if len(batch) < page_size:
            break

        page_no += 1

    logger.info(f"Successfully retrieved Insubiz vehicles. Total records: {len(all_rows)}")

    return all_rows


def fetch_insubiz_customers(
    http_hook: HttpHook,
    page_size: int = 400,
) -> list[dict[str, Any]]:
    """
    Fetch all customers from the Insubiz Cloud API (paged)

    This is used to build a lookup (customer_id -> customer dict) so that vehicles can
    be enriched with Level1..Level6 based on the customer parent hierarchy.

    :param http_hook: Airflow HttpHook configured for the Insubiz Cloud API connection.
    :param page_size: Page size used for pagination requests.
    :return: List of customer data
    """
    logger.info("Fetching customers from Insubiz Cloud API ...")

    jwt = _insubiz_sign_in(http_hook=http_hook)

    page_no = 1
    all_rows: list[dict] = []

    headers = {
        "Authorization": f"Bearer {jwt}",
        "Content-Type": "application/json",
        "Accept": "application/json",
    }

    while True:
        payload = {"pageNo": page_no, "pageSize": page_size}

        response = http_hook.run(
            endpoint="/api/v1.3/Customer/GetCustomersPagedAsync",
            json=payload,
            headers=headers,
        )

        response.raise_for_status()
        body: Any = response.json()

        if isinstance(body, dict):
            batch = body.get("data", []) or []
        elif isinstance(body, list):
            batch = body
        else:
            raise ValueError(f"Unexpected response type from customers endpoint: {type(body)}")

        if not isinstance(batch, list):
            raise ValueError("Expected a list of customers in response under key 'data'.")

        all_rows.extend(batch)

        logger.info(f"Fetched customers pageNo={page_no} rows={len(batch)} total={len(all_rows)}")

        if len(batch) < page_size:
            break

        page_no += 1

    logger.info(f"Successfully retrieved Insubiz customers. Total records: {len(all_rows)}")
    return all_rows


def update_insubiz_vehicle_fields(
    http_hook: HttpHook,
    record_id: int,
    fields: list[dict[str, Any]],
) -> dict[str, Any]:
    """
    Update fields on a single Insubiz vehicle.

    :param http_hook: Airflow HttpHook configured for the Insubiz Cloud API connection.
    :param record_id: Insubiz vehicle record ID to update.
    :param fields: List of field update objects expected by Insubiz API.
    :return: Response payload from Insubiz as a dictionary.
    """
    logger.info("Updating Insubiz vehicle fields ...")

    jwt = _insubiz_sign_in(http_hook=http_hook)

    headers = {
        "Authorization": f"Bearer {jwt}",
        "Content-Type": "application/json",
        "Accept": "application/json",
    }

    payload = {
        "recordId": record_id,
        "fields": fields,
    }

    logger.info(f"Payload: {payload}")

    response = http_hook.run(
        endpoint="/api/v1.3/Vehicle/UpdateVehicleFieldsAsync",
        json=payload,
        headers=headers,
    )

    response.raise_for_status()

    body: Any = response.json()
    if not isinstance(body, dict):
        raise ValueError(
            f"Unexpected update response type: {type(body)}")

    return body


def close_insubiz_vehicles_by_ids(
    http_hook: HttpHook,
    vehicle_ids: list[int],
    end_date_utc: str | None = None,
) -> int:
    """
    Set endDate for all given vehicle IDs (soft delete).

    :param http_hook: Airflow HttpHook configured for the Insubiz Cloud API connection.
    :param vehicle_ids: List of Insubiz vehicle record IDs to close.
    :param end_date_utc: Optional UTC timestamp string to set as ``endDate``.
    :return: Number of vehicles updated.
    """
    if not vehicle_ids:
        logger.info("No vehicle IDs to close.")
        return 0

    if not end_date_utc:
        end_date_utc = datetime.now(timezone.utc).isoformat(timespec="milliseconds").replace("+00:00", "Z")

    updated = 0
    for vehicle_id in vehicle_ids:
        update_insubiz_vehicle_fields(
            http_hook=http_hook,
            record_id=vehicle_id,
            fields=[
                {
                    "name": "endDate",
                    "value": end_date_utc,
                }
            ],
        )
        updated += 1

    logger.info(f"Updated endDate on {updated} vehicle(s).")
    return updated


def read_vehicle_ids_to_delete_from_excel_bytes(
    excel_bytes: bytes,
    sheet_name: str = "Skal slettes",
    id_column: str = "id",
) -> list[int]:
    """
    Read vehicle IDs from the delete sheet in the Excel file.

    :param excel_bytes: Excel file content as bytes (.xlsx).
    :param sheet_name: Sheet name containing vehicles to close.
    :param id_column: Column name containing Insubiz vehicle IDs.
    :return: Ordered list of unique vehicle IDs.
    """
    df = pd.read_excel(
        io.BytesIO(excel_bytes),
        engine="openpyxl",
        sheet_name=sheet_name,
    )

    normalized_cols = {str(c).strip().lower(): c for c in df.columns}
    source_col = normalized_cols.get(id_column.lower())

    if not source_col:
        raise ValueError(f"Missing required column '{id_column}' in sheet '{sheet_name}'")

    id_series = pd.to_numeric(df[source_col], errors="coerce").dropna().astype(int)

    # Preserve order, remove duplicates
    seen: set[int] = set()
    ids: list[int] = []
    for value in id_series.tolist():
        if value not in seen:
            seen.add(value)
            ids.append(value)

    return ids


def dfs_to_excel_bytes(sheets: dict[str, pd.DataFrame]) -> bytes:
    """
    Write one or more DataFrames to an Excel workbook and return the bytes.

    :param sheets: Mapping of sheet_name -> DataFrame.
    :return: Excel file content as bytes (.xlsx).
    """
    output = io.BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            sheet_name = name[:31]  # Excel sheet name limit
            frame.to_excel(writer, index=False, sheet_name=sheet_name)

            ws = writer.sheets[sheet_name]

            for idx, col in enumerate(frame.columns, start=1):
                max_length = max(frame[col].astype(str).map(len).max(), len(str(col)))
                ws.column_dimensions[get_column_letter(idx)].width = min(max_length + 2, 50) # max width of 50

    return output.getvalue()


def _norm_header(value: str) -> str:
    """
    Normalize a header value for tolerant matching.

    :param value: Header text to normalize.
    :return: Normalized header string.
    """
    return re.sub(r"\s+", " ", (value or "").strip().lower())


def read_motorstyrelsen_pdf_bytes(pdf_bytes: bytes) -> pd.DataFrame:
    """
    Parse Motorstyrelsen PDF bytes and extract vehicle table data.

    :param pdf_bytes: Raw PDF content as bytes.
    :return: DataFrame containing MOTORSTYRELSEN_PDF_COLUMNS plus "registreringsnummer".
    """
    expected_norm = {_norm_header(value=c): c for c in MOTORSTYRELSEN_PDF_COLUMNS}
    frames: list[pd.DataFrame] = []

    def _promote_first_row_as_header(frame: pd.DataFrame) -> pd.DataFrame:
        if frame.empty:
            return frame
        new_cols = [str(v).strip() for v in frame.iloc[0].tolist()]
        out = frame.iloc[1:].copy()
        out.columns = new_cols
        return out.reset_index(drop=True)

    def _header_match_count(cols: list[str]) -> int:
        return sum(1 for c in cols if _norm_header(value=c) in expected_norm)

    with fitz.open(stream=pdf_bytes, filetype="pdf") as doc:
        for page in doc:
            table_finder = page.find_tables()
            if not table_finder or not table_finder.tables:
                continue

            for table in table_finder.tables:
                table_df = table.to_pandas()
                if table_df is None or table_df.empty:
                    continue

                current_cols = [str(c).strip() for c in table_df.columns]
                col_matches = _header_match_count(cols=current_cols)

                # If headers are not detected in column names, try first row as header.
                # if col_matches < 3 and len(table_df) > 0:
                if col_matches < 3:
                    promoted = _promote_first_row_as_header(frame=table_df)
                    promoted_cols = [str(c).strip() for c in promoted.columns]
                    promoted_matches = _header_match_count(cols=promoted_cols)
                    if promoted_matches > col_matches:
                        table_df = promoted
                        current_cols = promoted_cols

                rename_map: dict[str, str] = {}
                for col in current_cols:
                    norm = _norm_header(value=col)
                    if norm in expected_norm:
                        rename_map[col] = expected_norm[norm]

                table_df = table_df.rename(columns=rename_map)

                keep_cols = [c for c in MOTORSTYRELSEN_PDF_COLUMNS if c in table_df.columns]
                if not keep_cols:
                    continue

                frames.append(table_df[keep_cols].copy())

    if not frames:
        raise ValueError("No tabular data found in Motorstyrelsen PDF")

    motor_df = pd.concat(frames, ignore_index=True)

    for col in MOTORSTYRELSEN_PDF_COLUMNS:
        if col not in motor_df.columns:
            motor_df[col] = pd.NA

    reg = (
        motor_df["Nummer (DPRN)"]
        .astype("string")
        .str.strip()
        .str.upper()
    )

    reg = reg.where(reg.str.fullmatch(r"[A-Z0-9]{1,7}", na=False))

    motor_df["registreringsnummer"] = reg
    motor_df = motor_df.loc[motor_df["registreringsnummer"].notna()].copy()

    if motor_df.empty:
        raise ValueError("PDF parse produced zero valid registreringsnummer")

    return motor_df[MOTORSTYRELSEN_PDF_COLUMNS + ["registreringsnummer"]]


def _norm_reg(series: pd.Series) -> pd.Series:
    """
    Normalize registration number values for stable comparisons.

    :param series: Input series containing registration numbers.
    :return: Normalized series.
    """
    return series.astype(str).str.strip().str.upper()


def compare_motorstyrelsen_vs_insubiz(
    motor_df: pd.DataFrame,
    insubiz_df: pd.DataFrame,
    motor_reg_col: str = "registreringsnummer",
    insubiz_reg_col: str = "Reg.nr.",
    exclude_prefix: str = "UINDREG",
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """
    Compare Motorstyrelsen vs Insubiz vehicles by registration number.

    Output:
    - need_to_delete: vehicles present in Insubiz but not in Motorstyrelsen
    - need_to_add: vehicles present in Motorstyrelsen but not in Insubiz

    :param motor_df: Motorstyrelsen DataFrame.
    :param insubiz_df: Insubiz DataFrame.
    :param motor_reg_col: Column name in motor_df containing registration numbers.
    :param insubiz_reg_col: Column name in insubiz_df containing registration numbers.
    :param exclude_prefix: Exclude registration numbers starting with this prefix.
    :return: (need_to_delete_df, need_to_add_df)
    """
    motor_key = _norm_reg(series=motor_df[motor_reg_col])
    insubiz_key = _norm_reg(series=insubiz_df[insubiz_reg_col])

    # Only keep valid Reg.nr that contains at most 7 alphanumeric characters
    insubiz_valid = insubiz_key.str.fullmatch(
        r"^[A-Z0-9]{1,7}$",
        na=False,
    )

    motor_prefix_ok = ~motor_key.str.startswith(pat=exclude_prefix, na=False)
    insubiz_prefix_ok = ~insubiz_key.str.startswith(pat=exclude_prefix, na=False)

    motor_mask = motor_prefix_ok
    insubiz_mask = insubiz_valid & insubiz_prefix_ok

    motor_df2 = motor_df.loc[motor_mask].copy()
    insubiz_df2 = insubiz_df.loc[insubiz_mask].copy()

    motor_key2 = _norm_reg(series=motor_df2[motor_reg_col])
    insubiz_key2 = _norm_reg(series=insubiz_df2[insubiz_reg_col])

    need_to_add = motor_df2.loc[~motor_key2.isin(set(insubiz_key2))].copy()
    logger.info(f"Found {len(need_to_add)} vehicles that needs to be added (in Insubiz)")

    need_to_delete = insubiz_df2.loc[~insubiz_key2.isin(set(motor_key2))].copy()
    logger.info(f"Found {len(need_to_delete)} vehicles that needs to be deleted (in Insubiz)")

    return need_to_delete, need_to_add


def _to_valid_id_or_none(value: Any) -> int | None:
    """
    Convert a value to int if possible, otherwise return None.

    :param value: Any value representing an id (int/str/None).
    :return: Integer value, or None if conversion fails or value is empty/zero.
    """
    try:
        if value is None:
            return None
        v = int(value)
        return v if v != 0 else None
    except (TypeError, ValueError):
        return None


def build_customer_levels(
    customer_id: Any,
    customer_by_id: dict[int, dict],
    root_id: int = INSUBIZ_ROOT_CUSTOMER_ID,
    max_levels: int = INSUBIZ_MAX_LEVELS,
) -> dict[str, str]:
    """
    Build Level1..LevelN labels for a customer based on parent hierarchy.

    :param customer_id: Customer ID for which levels should be generated.
    :param customer_by_id: Mapping of customer id -> customer payload.
    :param root_id: Root customer ID where the level path should start.
    :param max_levels: Maximum number of level columns to populate.
    :return: Dictionary with keys ``Level1``..``Level{max_levels}``.
    """
    levels = {f"Level{i}": "" for i in range(1, max_levels + 1)}

    current_id = _to_valid_id_or_none(value=customer_id)
    if not current_id:
        return levels

    seen: set[int] = set()
    chain_leaf_to_root: list[dict] = []

    while current_id and current_id not in seen:
        seen.add(current_id)

        customer = customer_by_id.get(current_id)
        if not isinstance(customer, dict):
            break

        chain_leaf_to_root.append(customer)

        parent = customer.get("parent") or {}
        parent_id = _to_valid_id_or_none(value=parent.get("id"))
        if not parent_id:
            break

        current_id = parent_id

    chain_root_to_leaf = list(reversed(chain_leaf_to_root))

    root_index = next(
        (
            i
            for i, c in enumerate(chain_root_to_leaf)
            if _to_valid_id_or_none(value=c.get("id")) == root_id
        ),
        None,
    )
    if root_index is None:
        return levels

    path = chain_root_to_leaf[root_index:]
    for i, node in enumerate(path[:max_levels], start=1):
        levels[f"Level{i}"] = (node.get("name") or "").strip()

    return levels


def enrich_vehicles_with_customer_levels(
    vehicles_df: pd.DataFrame,
    customers: list[dict[str, Any]],
) -> pd.DataFrame:
    """
    Enrich vehicle rows with organizational levels derived from customers.

    :param vehicles_df: Vehicle DataFrame to enrich.
    :param customers: List of customer payloads from Insubiz.
    :return: Vehicle DataFrame with Level1..Level6 columns.
    """
    df = vehicles_df.copy()

    customer_by_id = {
        int(c["id"]): c
        for c in customers
        if isinstance(c, dict) and c.get("id") is not None
    }

    if "customer_id" in df.columns:
        df["customer_id"] = pd.to_numeric(df["customer_id"], errors="coerce").astype("Int64")
        unique_customer_ids = df["customer_id"].dropna().astype(int).unique().tolist()

        levels_map = {
            cid: build_customer_levels(customer_id=cid, customer_by_id=customer_by_id)
            for cid in unique_customer_ids
        }

        levels_df = pd.DataFrame.from_dict(levels_map, orient="index").reset_index(names="customer_id")
        levels_df["customer_id"] = levels_df["customer_id"].astype("Int64")

        df = df.merge(levels_df, on="customer_id", how="left")
    else:
        for i in range(1, 7):
            df[f"Level{i}"] = ""

    return df


def find_latest_attachment(
    email_reader: EmailReader,
    mailbox: str = "INBOX",
    criteria: str = "UNSEEN",
    extensions: Sequence[str] = (".pdf",),
    filename_prefixes: Iterable[str] | None = None,
    max_emails: int = 50,
) -> tuple[bytes, str, bytes] | None:
    """
    Find newest attachment matching extension and filename prefixes.

    :param email_reader: EmailReader instance used to fetch emails.
    :param mailbox: Mailbox/folder name to search in.
    :param criteria: IMAP search criteria (for example ``UNSEEN`` or ``ALL``).
    :param extensions: Allowed file extensions (case-insensitive).
    :param filename_prefixes: Optional allowed filename prefixes.
    :param max_emails: Maximum number of emails to inspect.
    :return: Tuple ``(uid, filename, content_bytes)`` or ``None`` if no match.
    """
    emails, failed = email_reader.get_emails(
        mailbox=mailbox,
        criteria=criteria,
        set_flags=None,
        max=max_emails,
        low_to_high=False,
    )

    logger.info(f"Fetched {len(emails)} email(s), {len(failed)} failed to fetch.")

    extensions = tuple(ext.lower() for ext in extensions)
    prefixes = (
        tuple(p.lower() for p in filename_prefixes)
        if filename_prefixes
        else None
    )

    for msg in emails:
        uid: bytes = getattr(msg, "uid", None)

        for part in msg.iter_attachments():
            filename = part.get_filename() or ""
            filename_lc = filename.lower()

            if not filename_lc.endswith(extensions):
                continue

            if prefixes and not any(
                filename_lc.startswith(prefix)
                for prefix in prefixes
            ):
                continue

            content = part.get_payload(decode=True)
            if content:
                return uid, filename, content

    return None


def create_insubiz_vehicle(
    http_hook: HttpHook,
    payload: dict[str, Any],
) -> dict[str, Any]:
    """
    Create a single vehicle in Insubiz.

    :param http_hook: Airflow HttpHook configured for the Insubiz Cloud API connection.
    :param payload: Request payload for ``CreateVehicleAsync``.
    :return: Response payload from Insubiz as a dictionary.
    :raises ValueError: If API response payload is not a dictionary.
    """
    logger.info("Creating Insubiz vehicle ...")

    jwt = _insubiz_sign_in(http_hook=http_hook)

    headers = {
        "Authorization": f"Bearer {jwt}",
        "Content-Type": "application/json",
        "Accept": "application/json",
    }

    response = http_hook.run(
        endpoint="/api/v1.3/Vehicle/CreateVehicleAsync",
        json=payload,
        headers=headers,
    )
    response.raise_for_status()

    logger.info(f"Payload: {payload}")

    body: Any = response.json()
    if not isinstance(body, dict):
        raise ValueError(f"Unexpected create response type: {type(body)}")

    return body


def create_insubiz_vehicles_by_payloads(
    http_hook: HttpHook,
    payloads: list[dict[str, Any]],
) -> int:
    """
    Create multiple Insubiz vehicles from a list of payloads.

    :param http_hook: Airflow HttpHook configured for the Insubiz Cloud API connection.
    :param payloads: List of payloads for vehicle creation.
    :return: Number of vehicles created.
    """
    if not payloads:
        logger.info("No vehicles to create.")
        return 0

    created = 0
    for payload in payloads:
        create_insubiz_vehicle(http_hook=http_hook, payload=payload)
        created += 1

    logger.debug(f"Created {created} vehicle(s) in Insubiz.")
    return created


def read_vehicles_to_add_from_excel_bytes(
    excel_bytes: bytes,
    sheet_name: str = "Skal tilføjes",
    customer_id_column: str = "Customer_ID",
    registration_column: str = "Nummer (DPRN)",
    make_column: str = "Mærke beskrivelse (DKKØ)",
    chassis_column: str = "Stelnummer (DKKØ)",
    vehicle_art_column: str = "Køretøj art beskrivelse (DKRG)",
    model_description_column: str = "Model beskrivelse (DKKØ)",
    variant_description_column: str = "Variant beskrivelse (DKKØ)",
    registration_date_column: str = "Dato - første registrering (DKKØ)",
    driver_regular_column: str = "Bruger p-nr navn",
    model_year_column: str = "Model år (DKKØ)",
) -> list[dict[str, Any]]:
    """
    Read vehicle rows from Excel and convert them to Insubiz create payloads.

    :param excel_bytes: Excel file content as bytes (.xlsx).
    :param sheet_name: Sheet containing vehicles to create.
    :param customer_id_column: Column name for customer ID.
    :param registration_column: Column name for registration number.
    :param make_column: Column name for vehicle make.
    :param chassis_column: Column name for chassis number.
    :param vehicle_art_column: Column name for usage/category text.
    :param model_description_column: Column name for model description.
    :param variant_description_column: Column name for variant description.
    :param registration_date_column: Column name for first registration date.
    :param driver_regular_column: Column name for regular driver value.
    :param model_year_column: Column name for model year.
    :return: List of payload dictionaries for Insubiz vehicle creation.
    :raises ValueError: If required columns are missing or row values are
        invalid.
    """
    df = pd.read_excel(
        io.BytesIO(excel_bytes),
        engine="openpyxl",
        sheet_name=sheet_name,
    )

    normalized_cols = {str(c).strip().lower(): c for c in df.columns}

    logger.debug("Columns raw: %s", [repr(c) for c in df.columns])

    customer_col = normalized_cols.get(customer_id_column.strip().lower())
    registration_col = normalized_cols.get(registration_column.strip().lower())
    make_col = normalized_cols.get(make_column.strip().lower())
    chassis_col = normalized_cols.get(chassis_column.strip().lower())
    vehicle_type_col = normalized_cols.get(vehicle_art_column.strip().lower())
    model_description_col = normalized_cols.get(model_description_column.strip().lower())
    variant_description_col = normalized_cols.get(variant_description_column.strip().lower())
    registration_date_col = normalized_cols.get(registration_date_column.strip().lower())
    driver_regular_col = normalized_cols.get(driver_regular_column.strip().lower())
    model_year_col = normalized_cols.get(model_year_column.strip().lower())

    missing = []
    if not customer_col:
        missing.append(customer_id_column)
    if not registration_col:
        missing.append(registration_column)
    if not make_col:
        missing.append(make_column)
    if not chassis_col:
        missing.append(chassis_column)
    if not vehicle_type_col:
        missing.append(vehicle_art_column)
    if not model_description_col:
        missing.append(model_description_column)
    if not variant_description_col:
        missing.append(variant_description_column)
    if not registration_date_col:
        missing.append(registration_date_column)
    if not driver_regular_col:
        missing.append(driver_regular_column)
    if not model_year_col:
        missing.append(model_year_column)

    if missing:
        raise ValueError(
            f"Missing required column(s) in sheet '{sheet_name}': {', '.join(missing)}"
        )

    usage_text_to_id_map = {
        text.strip().lower(): int(code)
        for code, text in USAGE_TEXT_TO_CATEGORY_MAP.items()
    }

    start_date_utc = datetime.now(timezone.utc).isoformat(
        timespec="milliseconds"
    ).replace("+00:00", "Z")

    payloads: list[dict[str, Any]] = []

    for idx, row in df.iterrows():
        excel_row = idx + 2

        customer_raw = row.get(customer_col)
        customer_id_num = pd.to_numeric(customer_raw, errors="coerce")

        if pd.isna(customer_id_num):
            raise ValueError(
                f"Invalid or missing Customer_ID in sheet '{sheet_name}' at row {excel_row}"
            )

        customer_id = int(customer_id_num)
        if customer_id <= 0:
            raise ValueError(
                f"Customer_ID must be > 0 in sheet '{sheet_name}' at row {excel_row}"
            )

        registration = str(row.get(registration_col) or "").strip()
        make = str(row.get(make_col) or "").strip()
        chassis = str(row.get(chassis_col) or "").strip()
        usage_raw = str(row.get(vehicle_type_col) or "").strip()
        driver_regular = str(row.get(driver_regular_col) or "").strip()

        model_description = str(row.get(model_description_col) or "").strip()
        variant_description = str(row.get(variant_description_col) or "").strip()

        if not model_description:
            raise ValueError(
                f"Missing '{model_description_column}' in sheet '{sheet_name}' at row {excel_row}"
            )

        if not variant_description:
            raise ValueError(
                f"Missing '{variant_description_column}' in sheet '{sheet_name}' at row {excel_row}"
            )

        model = f"{model_description} - {variant_description}"

        registration_date_raw = row.get(registration_date_col)
        registration_date_dt = pd.to_datetime(
            registration_date_raw,
            format="%d-%m-%Y",
            errors="coerce",
        )

        if pd.isna(registration_date_dt):
            raise ValueError(
                f"Invalid or missing '{registration_date_column}' in sheet '{sheet_name}' at row {excel_row}"
            )

        if registration_date_dt.tzinfo is None:
            registration_date_dt = registration_date_dt.tz_localize("UTC")
        else:
            registration_date_dt = registration_date_dt.tz_convert("UTC")

        registration_date_utc = registration_date_dt.isoformat(
            timespec="milliseconds"
        ).replace("+00:00", "Z")

        model_year_raw = row.get(model_year_col)
        model_year_text = str(model_year_raw or "").strip()

        # Handle edge case when model year is represented as "-"
        if model_year_text == "-":
            model_year = 0
        else:
            model_year_num = pd.to_numeric(model_year_raw, errors="coerce")
            if pd.isna(model_year_num):
                raise ValueError(
                    f"Invalid or missing '{model_year_column}' in sheet '{sheet_name}' at row {excel_row}"
                )
            model_year = int(model_year_num)

        if not usage_raw:
            raise ValueError(
                f"Missing '{vehicle_art_column}' in sheet '{sheet_name}' at row {excel_row}"
            )

        usage_key = usage_raw.strip().lower()
        usage_id = usage_text_to_id_map.get(usage_key)

        if usage_id is None:
            raise ValueError(
                f"Unknown usage text '{usage_raw}' in sheet '{sheet_name}' at row {excel_row}"
            )

        usage_text = USAGE_TEXT_TO_CATEGORY_MAP[str(usage_id)]

        payload: dict[str, Any] = {
            "customer": {"id": customer_id},
            "driverRegular": driver_regular,
            "registration": registration,
            "registrationDate": registration_date_utc,
            "startDate": start_date_utc,
            "make": make,
            "model": model,
            "modelYear": model_year,
            "chassisNumber": chassis,
            "usage": {
                "text": usage_text,
                "id": usage_id,
            },
        }

        payloads.append(payload)

    return payloads
